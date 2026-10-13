#!/usr/bin/env python3
"""
Generate a PowerPoint presentation for the Chat Sphere project.
Theme matches Anurag University template: red/navy diagonals, Montserrat font.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import datetime
import math

# ─── Theme Colors ───
RED        = RGBColor(0xC0, 0x00, 0x00)
DARK_RED   = RGBColor(0xA0, 0x00, 0x00)
NAVY       = RGBColor(0x1B, 0x36, 0x5D)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BLACK      = RGBColor(0x00, 0x00, 0x00)
DARK_TEXT   = RGBColor(0x2D, 0x2D, 0x2D)
GRAY_TEXT   = RGBColor(0x66, 0x66, 0x66)
LIGHT_GRAY  = RGBColor(0xAA, 0xAA, 0xAA)
BG_LIGHT    = RGBColor(0xF8, 0xF8, 0xF8)
ACCENT_BLUE = RGBColor(0x00, 0x56, 0xA0)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

FONT_NAME = "Montserrat"


def add_freeform_triangle(slide, vertices, fill_color, line_color=None):
    """Add a freeform polygon (triangle) to the slide using oxml."""
    sp_tree = slide.shapes._spTree
    # Calculate bounding box
    xs = [v[0] for v in vertices]
    ys = [v[1] for v in vertices]
    left = min(xs)
    top = min(ys)
    width = max(xs) - left
    height = max(ys) - top

    sp = sp_tree.makeelement(qn('p:sp'), {})
    sp_tree.append(sp)

    nvSpPr = sp.makeelement(qn('p:nvSpPr'), {})
    sp.append(nvSpPr)
    cNvPr = nvSpPr.makeelement(qn('p:cNvPr'), {'id': '0', 'name': 'Freeform'})
    nvSpPr.append(cNvPr)
    cNvSpPr = nvSpPr.makeelement(qn('p:cNvSpPr'), {})
    nvSpPr.append(cNvSpPr)
    nvPr = nvSpPr.makeelement(qn('p:nvPr'), {})
    nvSpPr.append(nvPr)

    spPr = sp.makeelement(qn('p:spPr'), {})
    sp.append(spPr)

    xfrm = spPr.makeelement(qn('a:xfrm'), {})
    spPr.append(xfrm)
    off = xfrm.makeelement(qn('a:off'), {'x': str(int(left)), 'y': str(int(top))})
    xfrm.append(off)
    ext = xfrm.makeelement(qn('a:ext'), {'cx': str(int(width)), 'cy': str(int(height))})
    xfrm.append(ext)

    custGeom = spPr.makeelement(qn('a:custGeom'), {})
    spPr.append(custGeom)
    avLst = custGeom.makeelement(qn('a:avLst'), {})
    custGeom.append(avLst)
    gdLst = custGeom.makeelement(qn('a:gdLst'), {})
    custGeom.append(gdLst)
    ahLst = custGeom.makeelement(qn('a:ahLst'), {})
    custGeom.append(ahLst)
    cxnLst = custGeom.makeelement(qn('a:cxnLst'), {})
    custGeom.append(cxnLst)
    rect = custGeom.makeelement(qn('a:rect'), {'l': 'l', 't': 't', 'r': 'r', 'b': 'b'})
    custGeom.append(rect)

    pathLst = custGeom.makeelement(qn('a:pathLst'), {})
    custGeom.append(pathLst)
    path = pathLst.makeelement(qn('a:path'), {'w': str(int(width)), 'h': str(int(height))})
    pathLst.append(path)

    # Move to first vertex
    moveTo = path.makeelement(qn('a:moveTo'), {})
    path.append(moveTo)
    pt = moveTo.makeelement(qn('a:pt'), {
        'x': str(int(vertices[0][0] - left)),
        'y': str(int(vertices[0][1] - top))
    })
    moveTo.append(pt)

    # Line to remaining vertices
    for v in vertices[1:]:
        lnTo = path.makeelement(qn('a:lnTo'), {})
        path.append(lnTo)
        pt = lnTo.makeelement(qn('a:pt'), {
            'x': str(int(v[0] - left)),
            'y': str(int(v[1] - top))
        })
        lnTo.append(pt)

    close = path.makeelement(qn('a:close'), {})
    path.append(close)

    # Fill
    solidFill = spPr.makeelement(qn('a:solidFill'), {})
    spPr.append(solidFill)
    srgb = solidFill.makeelement(qn('a:srgbClr'), {
        'val': f'{fill_color[0]:02X}{fill_color[1]:02X}{fill_color[2]:02X}'
    })
    solidFill.append(srgb)

    # Line
    ln = spPr.makeelement(qn('a:ln'), {'w': '0'})
    spPr.append(ln)
    noFill = ln.makeelement(qn('a:noFill'), {})
    ln.append(noFill)

    return sp


def add_diagonal_header(slide):
    """Add the red/navy diagonal header decoration matching the Anurag theme."""
    W = int(SLIDE_W)
    H = int(SLIDE_H)

    # Large red triangle — top-left corner
    add_freeform_triangle(slide, [
        (0, 0),
        (Inches(3.2), 0),
        (0, Inches(1.8)),
    ], (0xC0, 0x00, 0x00))

    # Navy blue thin stripe — diagonal accent line below the red triangle
    add_freeform_triangle(slide, [
        (Inches(2.9), 0),
        (Inches(4.0), 0),
        (0, Inches(2.2)),
        (0, Inches(1.6)),
    ], (0x1B, 0x36, 0x5D))

    # Thin red diagonal line — top right area
    add_freeform_triangle(slide, [
        (Inches(10.0), 0),
        (Inches(13.333), 0),
        (Inches(13.333), Inches(0.12)),
        (Inches(9.8), Inches(0.12)),
    ], (0xC0, 0x00, 0x00))

    # Another thin navy line near top-right
    add_freeform_triangle(slide, [
        (Inches(10.5), Inches(0.18)),
        (Inches(13.333), Inches(0.18)),
        (Inches(13.333), Inches(0.28)),
        (Inches(10.3), Inches(0.28)),
    ], (0x1B, 0x36, 0x5D))


def add_university_text(slide):
    """Add Anurag University branding text (top-right)."""
    txBox = slide.shapes.add_textbox(Inches(9.5), Inches(0.35), Inches(3.6), Inches(0.55))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = "ANURAG UNIVERSITY"
    run.font.name = FONT_NAME
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = NAVY


def add_footer(slide, slide_number, total=11):
    """Add the footer bar: date | department | page number."""
    # Thin line above footer
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(6.95), Inches(12.333), Pt(1.5)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    line.line.fill.background()

    # Date
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(2), Inches(0.35))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = "7/7/2026"
    run.font.name = FONT_NAME
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY_TEXT

    # Department
    txBox = slide.shapes.add_textbox(Inches(4), Inches(7.05), Inches(5.333), Inches(0.35))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Department of Information Technology"
    run.font.name = FONT_NAME
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY_TEXT

    # Slide number
    txBox = slide.shapes.add_textbox(Inches(12.0), Inches(7.05), Inches(1), Inches(0.35))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = str(slide_number)
    run.font.name = FONT_NAME
    run.font.size = Pt(10)
    run.font.color.rgb = GRAY_TEXT


def add_slide_title(slide, title, subtitle=None):
    """Add a slide title with a left navy accent bar."""
    # Navy accent bar
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(0.08), Inches(0.55)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    # Title text
    txBox = slide.shapes.add_textbox(Inches(1.1), Inches(1.2), Inches(10), Inches(0.7))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.name = FONT_NAME
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = NAVY

    if subtitle:
        p2 = tf.add_paragraph()
        run2 = p2.add_run()
        run2.text = subtitle
        run2.font.name = FONT_NAME
        run2.font.size = Pt(14)
        run2.font.color.rgb = GRAY_TEXT


def add_body_text(slide, text, left=1.1, top=2.2, width=10.5, height=4.5, font_size=14, bold=False):
    """Add body text to the slide."""
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.space_after = Pt(6)
    p.line_spacing = Pt(22)
    run = p.add_run()
    run.text = text
    run.font.name = FONT_NAME
    run.font.size = Pt(font_size)
    run.font.color.rgb = DARK_TEXT
    run.font.bold = bold
    return tf


def add_bullet_points(slide, items, left=1.1, top=2.2, width=10.5, height=4.5, font_size=13):
    """Add bullet point items to the slide."""
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_after = Pt(6)
        p.space_before = Pt(3)
        p.line_spacing = Pt(20)
        p.level = 0

        # Red bullet marker
        bullet_run = p.add_run()
        bullet_run.text = "▸  "
        bullet_run.font.name = FONT_NAME
        bullet_run.font.size = Pt(font_size)
        bullet_run.font.color.rgb = RED
        bullet_run.font.bold = True

        run = p.add_run()
        run.text = item
        run.font.name = FONT_NAME
        run.font.size = Pt(font_size)
        run.font.color.rgb = DARK_TEXT

    return tf


def add_bordered_box(slide, left, top, width, height, header_text, body_lines, header_color=RED):
    """Add a box with left blue border, red header, and body text (like the title slide)."""
    # Left border accent
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.05), Inches(height)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()

    # Border rectangle (outline only)
    box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    box.fill.background()
    box.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
    box.line.width = Pt(0.75)

    # Header text
    txBox = slide.shapes.add_textbox(
        Inches(left + 0.2), Inches(top + 0.1), Inches(width - 0.3), Inches(0.3)
    )
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = header_text
    run.font.name = FONT_NAME
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = header_color

    # Body text
    y_offset = 0.35
    for line_text in body_lines:
        txBox2 = slide.shapes.add_textbox(
            Inches(left + 0.2), Inches(top + y_offset), Inches(width - 0.3), Inches(0.25)
        )
        tf2 = txBox2.text_frame
        p2 = tf2.paragraphs[0]
        run2 = p2.add_run()
        run2.text = line_text
        run2.font.name = FONT_NAME
        run2.font.size = Pt(13)
        run2.font.bold = "bold" in line_text.lower() or y_offset < 0.5
        run2.font.color.rgb = DARK_TEXT
        y_offset += 0.25


def make_table(slide, data, left, top, col_widths, font_size=11):
    """Create a styled table on the slide."""
    rows = len(data)
    cols = len(data[0])
    table_shape = slide.shapes.add_table(rows, cols, Inches(left), Inches(top),
                                          Inches(sum(col_widths)), Inches(rows * 0.4))
    table = table_shape.table

    # Set column widths
    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)

    for row_idx, row_data in enumerate(data):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = cell_text
            run.font.name = FONT_NAME
            run.font.size = Pt(font_size)

            if row_idx == 0:
                # Header row
                run.font.bold = True
                run.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            else:
                run.font.color.rgb = DARK_TEXT
                if row_idx % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xF0, 0xF4, 0xF8)
                else:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = WHITE

            # Vertical alignment
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE

    return table


# ═══════════════════════════════════════════════════════════════
# BUILD THE PRESENTATION
# ═══════════════════════════════════════════════════════════════

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slide_num = 0

# ─── SLIDE 1: TITLE ───
slide_num += 1
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
add_diagonal_header(slide)
add_university_text(slide)
add_footer(slide, slide_num)

# Main title
txBox = slide.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10), Inches(1.2))
tf = txBox.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run = p.add_run()
run.text = "Chat Sphere"
run.font.name = FONT_NAME
run.font.size = Pt(40)
run.font.bold = True
run.font.color.rgb = NAVY

# Subtitle
txBox2 = slide.shapes.add_textbox(Inches(2), Inches(3.0), Inches(9), Inches(0.6))
tf2 = txBox2.text_frame
p2 = tf2.paragraphs[0]
p2.alignment = PP_ALIGN.CENTER
run2 = p2.add_run()
run2.text = "A Full-Stack Web Application for Real time communication"
run2.font.name = FONT_NAME
run2.font.size = Pt(18)
run2.font.color.rgb = GRAY_TEXT

# Student Details box
add_bordered_box(slide, 2.5, 4.2, 3.8, 0.95,
                 "STUDENT DETAILS",
                 ["PurushothamReddy", "HT. No: 23EG112D51"])

# Project Guide box
add_bordered_box(slide, 7.0, 4.2, 3.8, 0.95,
                 "PROJECT GUIDE",
                 ["Dr. Niteesha Sharma, PhD", "HOD, IT Dept"])


# ─── SLIDE 2: ABSTRACT ───
slide_num += 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_diagonal_header(slide)
add_university_text(slide)
add_footer(slide, slide_num)
add_slide_title(slide, "Abstract")

abstract_text = (
    "Chat Sphere is a full-stack real-time messaging web application built on the "
    "MERN-style stack (React, Express, Node.js, and MongoDB) that delivers a comprehensive "
    "communication platform, enabling users to exchange instant text messages, share "
    "images and files via Cloudinary, create and manage group conversations, and track "
    "online presence and message delivery status — all from a single responsive, "
    "mobile-first web dashboard powered by Socket.IO for bidirectional real-time "
    "communication.\n\n"
    "The application implements secure JWT-based authentication, real-time typing "
    "indicators, read receipts with delivery status tracking (sent → delivered → read), "
    "group conversation management with admin controls, media sharing via Cloudinary CDN, "
    "voice message recording, in-chat message search, dark/light theme toggling, and "
    "a comprehensive notification system — deployed on Vercel (frontend), Render (backend), "
    "and MongoDB Atlas (database)."
)
add_body_text(slide, abstract_text, top=2.1, font_size=15)


# ─── SLIDE 3: PROBLEM STATEMENT ───
slide_num += 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_diagonal_header(slide)
add_university_text(slide)
add_footer(slide, slide_num)
add_slide_title(slide, "Problem Statement")

problems = [
    "Latency in Message Delivery — Traditional HTTP request-response cycles introduce unacceptable delays. Users expect sub-second delivery with real-time typing indicators and read receipts.",
    "Lack of Real-Time Presence Tracking — Users need to know whether contacts are currently online, last active time, and whether they are typing a response.",
    "Fragmented Group Communication — Creating group conversations, managing participants, and ensuring correct unread counts requires careful schema design.",
    "File and Media Sharing Complexity — Sharing images, documents, and voice messages seamlessly with cloud storage, size limits, and preview rendering.",
    "Security Concerns — User credentials must be hashed, tokens tamper-proof, APIs rate-limited, and messages encrypted at rest.",
]
add_bullet_points(slide, problems, top=2.1, font_size=13)


# ─── SLIDE 4: OBJECTIVES ───
slide_num += 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_diagonal_header(slide)
add_university_text(slide)
add_footer(slide, slide_num)
add_slide_title(slide, "Project Objectives")

objectives = [
    "Secure user registration and login using JWT with Bcrypt password hashing and token verification middleware.",
    "Real-time bidirectional messaging using Socket.IO with delivery acknowledgments and read receipts (✓, ✓✓, blue ✓✓).",
    "Real-time presence system tracking online/offline status and \"last seen\" timestamps via Socket.IO events.",
    "Typing indicators with animated feedback and automatic timeout-based clearing.",
    "Complete CRUD for conversations — direct and group chats with multi-step group creation modal.",
    "Rich message input — text, emoji, image/document attachments (Cloudinary, 10MB), voice recording, clipboard paste.",
    "Message management — deletion, in-chat search with match navigation, chat clearing, conversation deletion.",
    "Notifications panel, profile management, dark/light theme toggle, and cloud deployment (Vercel + Render + Atlas).",
]
add_bullet_points(slide, objectives, top=2.1, font_size=12)


# ─── SLIDE 5: ARCHITECTURE & TECH STACK ───
slide_num += 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_diagonal_header(slide)
add_university_text(slide)
add_footer(slide, slide_num)
add_slide_title(slide, "Architecture & Tech Stack")

arch_data = [
    ["Layer", "Components & Technologies"],
    ["Presentation", "React 19 + Vite 8, Tailwind CSS v3, Zustand, React Router v7, Axios, CryptoJS"],
    ["Application", "Node.js + Express.js, Socket.IO v4, JWT Auth, CORS, Helmet.js, Rate Limiting"],
    ["Data", "MongoDB Atlas (Cloud NoSQL), Mongoose ODM (User, Conversation, Message schemas)"],
    ["External Services", "Cloudinary CDN for image, file, and voice message uploads"],
    ["Deployment", "Frontend → Vercel  |  Backend → Render  |  Database → MongoDB Atlas"],
]
make_table(slide, arch_data, left=1.0, top=2.2, col_widths=[2.5, 8.5], font_size=12)


# ─── SLIDE 6: LITERATURE SURVEY ───
slide_num += 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_diagonal_header(slide)
add_university_text(slide)
add_footer(slide, slide_num)
add_slide_title(slide, "Literature Survey")

lit_data = [
    ["#", "Author(s) & Year", "Focus Area", "Key Finding / Limitation"],
    ["1", "Patel & Desai (2023)", "MERN + Socket.IO Chat", "Real-time messaging; lacks groups & media"],
    ["2", "Zhang et al. (2022)", "WebSocket Performance", "10x faster than HTTP polling; no typing/presence"],
    ["3", "Kumar & Singh (2024)", "JWT + bcrypt Security", "Strong auth; no real-time features"],
    ["4", "Rahman (2023)", "WebRTC Video Calling", "Standalone; not integrated into messaging"],
    ["5", "Sharma & Gupta (2025)", "Group Messaging (NoSQL)", "Efficient broadcast; lacks admin controls"],
    ["6", "Li & Wang (2022)", "Cloud File Sharing", "Seamless uploads; no real-time integration"],
    ["7", "Reddy & Rao (2024)", "Zustand vs Redux", "Simpler API; no messaging patterns"],
    ["8", "Fernandez (2023)", "Tailwind CSS UI", "Rapid dev; no dark mode / theme switching"],
    ["9", "Das & Mukherjee (2024)", "Read Receipts", "Accurate tracking; single-user demo only"],
    ["10", "Nakamura (2025)", "PWA Chat + Offline", "Installable as app; lacks advanced types"],
]
make_table(slide, lit_data, left=0.5, top=2.1, col_widths=[0.5, 2.5, 3.0, 5.5], font_size=10)


# ─── SLIDE 7: KEY MODULES ───
slide_num += 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_diagonal_header(slide)
add_university_text(slide)
add_footer(slide, slide_num)
add_slide_title(slide, "Key Modules & Features")

modules_data = [
    ["Module", "Description"],
    ["Authentication", "JWT + Bcrypt login/register; token verification middleware; session management"],
    ["Chat Dashboard", "Modern sidebar + chat window; responsive mobile-first layout"],
    ["Real-Time Messaging", "Socket.IO events for send/receive/delivered/read; optimistic UI updates"],
    ["Group Conversations", "Multi-step group creation; admin controls; add/remove members"],
    ["Media Sharing", "Cloudinary uploads (images, docs, voice); 10MB limit; clipboard paste"],
    ["Presence & Typing", "Online/offline status; last seen; animated typing indicators"],
    ["Notifications", "Slide-in panel with message previews, timestamps, read/unread status"],
    ["Profile Management", "Edit name, about, avatar; Cloudinary upload; user blocking"],
    ["Theme System", "Dark/light toggle with Zustand persistence; CSS custom properties"],
    ["Message Search", "In-chat search with match counter, navigation, and auto-scroll"],
]
make_table(slide, modules_data, left=0.8, top=2.1, col_widths=[2.5, 8.8], font_size=11)


# ─── SLIDE 8: METHODOLOGY ───
slide_num += 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_diagonal_header(slide)
add_university_text(slide)
add_footer(slide, slide_num)
add_slide_title(slide, "Methodology")

methodology = [
    "Requirement Analysis — Identified functional requirements (auth, messaging, groups, media, notifications) and non-functional requirements (latency, security, scalability).",
    "System Design — Layered architecture with 4 REST API route groups (/auth, /users, /conversations, /messages) and Socket.IO event contracts.",
    "Database Design — Three Mongoose collections: User, Conversation, Message with compound indexes for efficient queries.",
    "Backend Development — Express module-by-module: Auth routes, User routes, Conversation routes, Message routes, Socket event handlers, Middleware.",
    "Frontend Development — React component groups: Auth, Chat, Group, Profile, Common components + Zustand stores + Custom hooks.",
    "Integration — Vite dev-server proxy (local) + Vercel rewrite rules (production); Socket.IO with JWT handshake.",
    "Testing — Postman API testing + Manual UI testing of all interactions across mobile/tablet/desktop viewports.",
    "Deployment — Backend on Render, Frontend on Vercel, Database on MongoDB Atlas with environment variable configuration.",
]
add_bullet_points(slide, methodology, top=2.1, font_size=12)


# ─── SLIDE 9: DATABASE SCHEMA ───
slide_num += 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_diagonal_header(slide)
add_university_text(slide)
add_footer(slide, slide_num)
add_slide_title(slide, "Database Schema")

db_data = [
    ["Collection", "Key Fields", "Purpose"],
    ["User", "fullName, email, password, avatar,\nabout, lastSeen, isOnline, blockedUsers", "Account credentials, profile data,\nonline status, block list"],
    ["Conversation", "type, participants, groupInfo,\nlastMessage, unreadCount, wallpaper", "Links participants, stores lastMessage\n(subset pattern), per-user unread counts"],
    ["Message", "conversationId, sender, content, type,\nstatus, readBy, reactions, isDeleted", "Message logs with delivery tracking,\nread receipts, reactions, soft deletion"],
]
make_table(slide, db_data, left=0.8, top=2.2, col_widths=[2.0, 4.5, 4.8], font_size=11)

# Socket Events section
add_body_text(slide, "Socket.IO Event Contract:", left=0.8, top=4.8, width=10, height=0.4, font_size=14, bold=True)
socket_events = [
    "message:send / message:new / message:sent / message:delivered / message:read / message:delete",
    "typing:start / typing:stop  •  user:online / user:offline  •  conversation:new / group:updated",
]
add_bullet_points(slide, socket_events, left=0.8, top=5.2, width=11, height=1.5, font_size=12)


# ─── SLIDE 10: SOFTWARE & HARDWARE REQUIREMENTS ───
slide_num += 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_diagonal_header(slide)
add_university_text(slide)
add_footer(slide, slide_num)
add_slide_title(slide, "Software & Hardware Requirements")

sw_data = [
    ["Component", "Specification", "Purpose"],
    ["Runtime", "Node.js 18+", "Express backend + Vite build tooling"],
    ["Backend", "Express.js", "REST API, middleware, DB queries"],
    ["Frontend", "React 19 + Vite 8", "UI library + fast HMR dev server"],
    ["Real-Time", "Socket.IO v4", "WebSocket communication"],
    ["Database", "MongoDB Atlas + Mongoose", "Cloud NoSQL with schema validation"],
    ["Auth", "jsonwebtoken + bcryptjs", "JWT + password hashing (12 rounds)"],
    ["Media", "Cloudinary CDN", "Image, file, voice message uploads"],
    ["State Mgmt", "Zustand v5", "Lightweight React state management"],
    ["Styling", "Tailwind CSS v3, React Icons", "Utility-first CSS + icon library"],
    ["Security", "Helmet.js, rate-limit, CryptoJS", "Headers, rate limiting, AES encryption"],
]
make_table(slide, sw_data, left=0.5, top=2.1, col_widths=[2.0, 4.0, 5.5], font_size=10)

# Hardware requirements (below)
hw_text = "Hardware: Intel i3+, 4GB RAM (8GB recommended), 2GB storage, stable internet, microphone (for voice recording)"
add_body_text(slide, hw_text, left=0.8, top=6.3, width=11, height=0.5, font_size=11)


# ─── SLIDE 11: TIMELINE & REFERENCES ───
slide_num += 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_diagonal_header(slide)
add_university_text(slide)
add_footer(slide, slide_num)
add_slide_title(slide, "Project Timeline & References")

timeline_data = [
    ["Phase", "Month 1", "Month 2", "Month 3", "Month 4"],
    ["Requirement Analysis & Survey", "✓", "", "", ""],
    ["System Design & DB Schema", "✓", "", "", ""],
    ["Backend REST API Development", "", "✓", "✓", ""],
    ["Socket.IO Event Handlers", "", "✓", "✓", ""],
    ["Frontend Components", "", "✓", "✓", ""],
    ["Integration & Testing", "", "", "✓", "✓"],
    ["Cloud Deployment", "", "", "", "✓"],
    ["Documentation & Reporting", "", "", "", "✓"],
]
make_table(slide, timeline_data, left=0.8, top=2.1, col_widths=[4.0, 1.5, 1.5, 1.5, 1.5], font_size=11)

# Key references
refs_items = [
    "Patel & Desai (2023) — MERN + Socket.IO Chat  •  Zhang et al. (2022) — WebSocket Performance",
    "Kumar & Singh (2024) — JWT Security  •  Sharma & Gupta (2025) — Group Messaging Architecture",
    "MongoDB Atlas Docs  •  React Docs  •  Express.js Docs  •  Socket.IO Docs  •  Cloudinary Docs",
]
add_body_text(slide, "Key References:", left=0.8, top=5.6, width=10, height=0.3, font_size=14, bold=True)
add_bullet_points(slide, refs_items, left=0.8, top=5.9, width=11, height=1.0, font_size=11)


# ─── SAVE ───
output_path = "/Users/purushothamreddy/Downloads/extras/Projects/Chat_Application/Chat_Sphere_PPT.pptx"
prs.save(output_path)
print(f"✅ Presentation saved to: {output_path}")
print(f"   Total slides: {slide_num}")
